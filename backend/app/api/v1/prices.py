import io
import math
from datetime import date

from fastapi import APIRouter, Depends, Query, UploadFile, File
from sqlalchemy import func, select, and_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.dependencies import get_current_user, require_role
from app.core.rate_limit import enforce_upload_rate_limit
from app.models.enums import UserRole
from app.core.exceptions import BadRequestException, NotFoundException
from app.models.price_entry import PriceEntry
from app.models.spare_part import SparePart
from app.models.user import User
from app.schemas.common import PaginatedResponse

router = APIRouter(prefix="/prices", tags=["Prices"])


@router.get("/", response_model=PaginatedResponse[dict])
async def list_prices(
    page: int = Query(1, ge=1, le=10000),
    page_size: int = Query(20, ge=1, le=100),
    spare_part_id: int | None = Query(None, description="Filter by spare part"),
    currency: str | None = Query(None, description="Filter by currency"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List price entries with pagination and filtering."""
    query = select(PriceEntry)
    count_query = select(func.count(PriceEntry.id))

    conditions = []
    if spare_part_id is not None:
        conditions.append(PriceEntry.spare_part_id == spare_part_id)
    if currency:
        conditions.append(PriceEntry.currency == currency)

    if conditions:
        combined = and_(*conditions)
        query = query.where(combined)
        count_query = count_query.where(combined)

    total_result = await db.execute(count_query)
    total = total_result.scalar() or 0

    offset = (page - 1) * page_size
    query = query.order_by(PriceEntry.created_at.desc()).offset(offset).limit(page_size)

    result = await db.execute(query)
    prices = result.scalars().all()

    return {
        "items": [_price_to_dict(p) for p in prices],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": math.ceil(total / page_size) if total > 0 else 0,
    }


@router.post("/", status_code=201)
async def create_price(
    data: dict,
    current_user: User = Depends(require_role(UserRole.OPERATIONS)),
    db: AsyncSession = Depends(get_db),
):
    """Create a new price entry (operations only)."""
    spare_part_id = data.get("spare_part_id")
    list_price = data.get("list_price")
    net_price = data.get("net_price")

    if not spare_part_id or list_price is None or net_price is None:
        raise BadRequestException("spare_part_id, list_price, and net_price are required")

    # Verify spare part exists
    part_result = await db.execute(
        select(SparePart).where(SparePart.id == spare_part_id)
    )
    if not part_result.scalar_one_or_none():
        raise NotFoundException("Yedek parca bulunamadi")

    # Parse optional date fields
    valid_from = _parse_date(data.get("valid_from"))
    valid_until = _parse_date(data.get("valid_until"))

    price = PriceEntry(
        spare_part_id=spare_part_id,
        list_price=float(list_price),
        discount_pct=float(data.get("discount_pct", 0.0)),
        net_price=float(net_price),
        currency=data.get("currency", "USD"),
        valid_from=valid_from,
        valid_until=valid_until,
        price_list_version=data.get("price_list_version"),
    )
    db.add(price)
    await db.flush()
    await db.refresh(price)

    return _price_to_dict(price)


@router.post(
    "/import",
    status_code=201,
    dependencies=[Depends(enforce_upload_rate_limit)],
)
async def import_prices(
    file: UploadFile = File(...),
    current_user: User = Depends(require_role(UserRole.OPERATIONS)),
    db: AsyncSession = Depends(get_db),
):
    """Import price entries from Excel (.xlsx) or CSV file."""
    if not file.filename:
        raise BadRequestException("Dosya saglanmadi")

    filename_lower = file.filename.lower()
    if not (filename_lower.endswith(".csv") or filename_lower.endswith(".xlsx")):
        raise BadRequestException("Yalnizca .csv ve .xlsx dosyalari desteklenmektedir")

    from app.core.config import settings as cfg
    from app.core.upload_utils import read_upload_file_limited
    content = await read_upload_file_limited(file, max_bytes=cfg.MAX_UPLOAD_SIZE_MB * 1024 * 1024)

    imported_count = 0
    skipped_count = 0
    errors = []

    try:
        if filename_lower.endswith(".csv"):
            import csv

            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader)
        else:
            try:
                import openpyxl
            except ImportError:
                raise BadRequestException(
                    "openpyxl is required for Excel imports. Install it with: pip install openpyxl"
                )
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append(dict(zip(headers, row)))

        for i, row in enumerate(rows, start=2):
            honeywell_code = str(row.get("honeywell_code") or "").strip()
            if not honeywell_code:
                skipped_count += 1
                continue

            # Look up spare part by code
            part_result = await db.execute(
                select(SparePart).where(SparePart.honeywell_code == honeywell_code)
            )
            part = part_result.scalar_one_or_none()
            if not part:
                errors.append(f"Row {i}: Part '{honeywell_code}' not found")
                skipped_count += 1
                continue

            try:
                list_price = float(row.get("list_price") or 0)
                net_price = float(row.get("net_price") or 0)
            except (ValueError, TypeError):
                errors.append(f"Row {i}: Invalid price values")
                skipped_count += 1
                continue

            price = PriceEntry(
                spare_part_id=part.id,
                list_price=list_price,
                discount_pct=float(row.get("discount_pct") or 0.0),
                net_price=net_price,
                currency=str(row.get("currency") or "USD").strip(),
                valid_from=_parse_date(row.get("valid_from")),
                valid_until=_parse_date(row.get("valid_until")),
                price_list_version=str(row.get("price_list_version") or "").strip() or None,
            )
            db.add(price)
            imported_count += 1

        await db.flush()

    except BadRequestException:
        raise
    except Exception as e:
        raise BadRequestException(f"Error processing file: {str(e)}")

    return {
        "message": "Icerik aktarimi tamamlandi",
        "imported": imported_count,
        "skipped": skipped_count,
        "errors": errors,
    }


def _parse_date(value) -> date | None:
    """Parse a date from string or date object."""
    if value is None:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None
    return None


def _price_to_dict(price: PriceEntry) -> dict:
    """Convert PriceEntry to a dictionary response."""
    return {
        "id": price.id,
        "spare_part_id": price.spare_part_id,
        "list_price": price.list_price,
        "discount_pct": price.discount_pct,
        "net_price": price.net_price,
        "currency": price.currency,
        "valid_from": price.valid_from.isoformat() if price.valid_from else None,
        "valid_until": price.valid_until.isoformat() if price.valid_until else None,
        "price_list_version": price.price_list_version,
        "created_at": price.created_at.isoformat() if price.created_at else None,
        "spare_part": {
            "id": price.spare_part.id,
            "honeywell_code": price.spare_part.honeywell_code,
            "name_en": price.spare_part.name_en,
        } if price.spare_part else None,
    }
